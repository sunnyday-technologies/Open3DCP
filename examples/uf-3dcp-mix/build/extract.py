#!/usr/bin/env python3
"""Reproduce source-linked UF rheology/age rows without guessing material classifications.

Complete common-basis ratios can normalize without kg/m3. This excerpt leaves
composition unconverted because blanks and mixed reporting bases do not establish
a complete inventory. source_cells.csv preserves the reported composition cells.
"""
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path

SOURCE_SHA256 = "369369bb2b19610f2ca9eced6cfc395f91e5cf81fec6e63c1fc919bcc1983316"
DOI = "10.5281/zenodo.6828947"
COLS = ["source_dataset", "is_3d_printed", "material_class", "w_b_ratio",
        "static_yield_stress_pa", "dynamic_yield_stress_pa", "plastic_viscosity_pa_s",
        "test_age_days", "compressive_strength_mpa", "doi", "source_citation", "provenance_notes"]
IX = dict(ref=4, binder=5, binder1=6, wb=40, sy=97, dy=98, pv=99)
AGES = [(1, 100), (3, 101), (7, 102), (28, 103)]
MAX_SOURCE_ROWS = 10
MAX_PER_REF = 2
EXPECTED = {4: "Reference", 5: "Binder", 6: "Binder1", 40: "Water-Binder Ratio",
            97: "Static Yield Stress (kPa)", 98: "Dynamic Yield Stress (kPa)",
            99: "Plastic Viscosity", 100: "1days(MPa)", 101: "3days(MPa)",
            102: "7days(MPa)", 103: "28days(MPa)"}

def num(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        n = float(value)
        return n if math.isfinite(n) and n >= 0 else None
    except (ValueError, TypeError):
        return None

def text(value):
    return "" if value is None else str(value).strip()

def check_headers(headers):
    for i, expected in EXPECTED.items():
        if i >= len(headers) or expected.lower() not in " ".join(text(headers[i]).split()).lower():
            raise ValueError(f"UF header or unit changed at column {i}; review mapping")
    # The source encodes the Pa.s separator with a replacement character.
    if text(headers[99]) not in {"Plastic Viscosity (Pa.s)", "Plastic Viscosity (Pa·s)", "Plastic Viscosity (Pa�s)"}:
        raise ValueError("Viscosity unit changed")

def convert(headers, records):
    check_headers(headers)
    out, provenance, cells, per_ref = [], [], [], {}
    for excel_row, r in enumerate(records, start=2):
        if len(r) < 104:
            raise ValueError("Truncated UF row")
        g = lambda key: num(r[IX[key]])
        ages = [(day, num(r[i])) for day, i in AGES if num(r[i]) is not None]
        ref = text(r[IX["ref"]])
        if g("sy") is None or g("wb") is None or not ages or not ref:
            continue
        if per_ref.get(ref, 0) >= MAX_PER_REF:
            continue
        per_ref[ref] = per_ref.get(ref, 0) + 1
        # Preserve labels and quantities separately, including blanks. No name cell is parsed as a dose.
        for col in range(3, 97):
            cells.append({"source_row": excel_row, "column_index_zero_based": col,
                          "source_header": text(headers[col]), "source_value": text(r[col]),
                          "missing": r[col] is None or text(r[col]) == ""})
        provenance.append({"source_row": excel_row, "primary_reference": ref,
                           "binder_description": text(r[5]), "binder1_label": text(r[6]),
                           "age_rows": len(ages),
                           "composition_status": "unconverted: complete inventory/common basis not established",
                           "material_class_status": "unclassified; raw source labels retained",
                           "specimen_printing_status": "not established for each mechanical result"})
        note = (f"Sheet1 source row {excel_row}; primary study: {ref}. "
                f"Reported binder: {text(r[5]) or 'not reported'}; Binder1: {text(r[6]) or 'not reported'}. "
                "Raw composition labels, quantities, units and blanks: source_cells.csv. "
                "Complete binder-relative ratios can normalize without absolute binder dosage; "
                "this excerpt does not establish complete inventory or a common basis for all components. "
                "No constituent mass-% or kg/m3 conversion. Blank is not zero. "
                "Material class and mechanical specimen printing status are not inferred.")
        for day, strength in ages:
            out.append({"source_dataset": "UF 3DCP Mix-Design Open Dataset", "is_3d_printed": None,
                        "material_class": None, "w_b_ratio": g("wb"),
                        "static_yield_stress_pa": g("sy") * 1000,
                        "dynamic_yield_stress_pa": g("dy") * 1000 if g("dy") is not None else None,
                        "plastic_viscosity_pa_s": g("pv"), "test_age_days": day,
                        "compressive_strength_mpa": strength, "doi": DOI,
                        "source_citation": "Gao, J.; Wang, Z.; Wang, C. (2023). 3D Printing Concrete Mix Design Open Dataset (v0.3). DOI " + DOI,
                        "provenance_notes": note})
        if len(provenance) == MAX_SOURCE_ROWS:
            break
    return out, provenance, cells

def write_csv(path, fields, rows):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)

def run(source, output):
    from openpyxl import load_workbook
    source, output = Path(source).resolve(), Path(output).resolve()
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError("UF source fingerprint differs from audited v0.3 workbook")
    book = load_workbook(source, read_only=True, data_only=True)
    try:
        it = book["Sheet1"].iter_rows(values_only=True)
        rows, provenance, cells = convert(next(it), it)
    finally:
        book.close()
    output.mkdir(parents=True, exist_ok=True)
    write_csv(output / "uf-3dcp-mix.open3dcp.csv", COLS, rows)
    write_csv(output / "provenance.csv", ["source_row", "primary_reference", "binder_description",
              "binder1_label", "age_rows", "composition_status", "material_class_status", "specimen_printing_status"], provenance)
    write_csv(output / "source_cells.csv", ["source_row", "column_index_zero_based", "source_header", "source_value", "missing"], cells)
    report = {"source_sha256": digest, "source_version_doi": "10.5281/zenodo.8070144",
              "rows_committed": len(rows), "selected_source_rows": len(provenance),
              "selection": "first 10 source rows with finite nonnegative static yield, w/b, strength and reference; at most two rows per reference; all available ages",
              "row_grain": "one source mixture at one reported test age",
              "unit_conversion": "static/dynamic yield stress kPa to Pa multiplied by 1000",
              "composition_conversion": "not performed; incomplete inventory and mixed basis remain explicit",
              "printability_or_OPC_filter": "none; no claim that selection establishes either"}
    (output / "extraction_report.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
        raise RuntimeError("Source changed during extraction")
    print(json.dumps(report))

if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("source")
    p.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parents[1])
    a = p.parse_args()
    run(a.source, a.output_dir)
